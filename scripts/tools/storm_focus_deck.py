"""
Focused briefing deck built from the same published figures as
overview_powerpoint.py (reuses its server/NHC/slide helpers):

  1. Eastern Pacific - Mexico: Hurricane Polo cold wake, Hurricane Rachel, and
     the NHC disturbance with a 90% 7-day formation chance. Temperature 0m and
     Ocean Heat Content only.
  2. maracoos_007 (Mid Atlantic Bight glider): daily glider profile figures
     plus Temperature 0m maps, Sep 21 -> today.
  3. Typhoon Choi-wan (Guam region): Ocean Heat Content only, from formation
     to today.

Usage:
    python storm_focus_deck.py
    python storm_focus_deck.py --output my_deck.pptx --hour 12Z
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import overview_powerpoint as o  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

EP = "Eastern Pacific - Mexico"
MAB = "Mid Atlantic Bight"
GUAM = "Guam"
GLIDER = "maracoos_007"

# Polo (EP17): closest approach 9/22 18Z, peak cooling 9/24 — see
# ~/Documents/polo_analysis/README.md
POLO_DAYS = [dt.date(2026, 9, d) for d in range(21, 28)]
# Rachel (EP18) formed 9/25; every other day plus today.
RACHEL_DAYS = [dt.date(2026, 9, 25), dt.date(2026, 9, 27), dt.date(2026, 9, 29),
               dt.date(2026, 10, 1), dt.date(2026, 10, 3)]
# Choi-wan (WP26) formed 9/30 at ~15.6N 151.9E.
CHOIWAN_DAYS = [dt.date(2026, 9, 30) + dt.timedelta(days=i) for i in range(6)]
GLIDER_START = dt.date(2026, 9, 21)


def section(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = prs.slide_width, prs.slide_height
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.4), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = o.NAVY
    bar.line.fill.background()
    o._text_box(slide, Inches(0.7), Inches(3.2), W - Inches(1), Inches(1.2),
                title, 40, bold=True, color=o.NAVY)
    o._text_box(slide, Inches(0.7), Inches(4.5), W - Inches(1), Inches(0.5),
                subtitle, 16, color=o.GREY)


def map_slides(prs, region, vd, days, hour, tag):
    """One image per slide so slides can be flipped through. Slides are grouped
    by comparison (all days of RTOFS vs CMEMS, then all days of RTOFS vs ESPC)
    so paging forward shows change through time. OHC falls back to the legacy
    combined panel when the standard files are missing."""
    label = o._product_label(vd)
    found = {}  # day -> (hour, [(buf, name), ...])
    for d in days:
        h = o.get_available_hour_for_date(region, vd, d, preferred_hour=hour, require="rtofs")
        if not h:
            print(f"    {d} {vd}: not found")
            continue
        cop, espc, _, _ = o.build_map_urls(region, vd, d, h)
        pairs = [(o.fetch_image(cop), "RTOFS vs CMEMS"), (o.fetch_image(espc), "RTOFS vs ESPC")]
        if not any(b for b, _ in pairs) and vd == "ocean_heat_content":
            pairs = [(o.fetch_image(o.build_ohc_legacy_url(region, d, h)), "RTOFS / ESPC / CMEMS")]
        if any(b for b, _ in pairs):
            found[d] = (h, pairs)
        else:
            print(f"    {d} {h} {vd}: images unavailable")
    n = 0
    names = []
    for _, (_, pairs) in found.items():
        for _, nm in pairs:
            if nm not in names:
                names.append(nm)
    for nm in names:
        for d, (h, pairs) in found.items():
            for buf, pn in pairs:
                if pn == nm and buf:
                    o.add_image_slide(prs, region, f"{tag}  ·  {label}  ·  {nm}",
                                      [(buf, nm)], d.isoformat(), h)
                    n += 1
        print(f"    {vd} {nm}: {n} slides so far")
    return n


def nhc_slides(prs, query, keys, today):
    storm = o.match_storm(o.discover_active_storms(), query)
    if not storm:
        print(f"  NHC: no active storm matching {query}")
        return
    for url in o.get_storm_graphic_urls(storm["slot"], storm["advisory_id"], keys):
        buf = o.fetch_image(url)
        if buf:
            lbl = o.nhc_image_label(url)
            o.add_image_slide(prs, "NHC Advisory", f"{storm['name']}  ·  {lbl}",
                              [(buf, lbl)], today, "")


def glider_slides(prs, today, end):
    d, n = GLIDER_START, 0
    while d <= end:
        url = next((u for g, u in o.get_glider_files(d).items() if g.startswith(GLIDER)), None)
        buf = o.fetch_image(url)
        if buf:
            o.add_image_slide(prs, GLIDER, "Glider Profiles vs Models (400 m)",
                              [(buf, f"{GLIDER}  ·  {d.isoformat()}")], d.isoformat(), "")
            n += 1
            print(f"    {d} profile ok")
        else:
            print(f"    {d} profile: not found")
        d += dt.timedelta(days=1)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", "-o", default=str(Path.home() / "Documents/2026-10-05_storm_focus.pptx"))
    ap.add_argument("--hour", default="12Z")
    args = ap.parse_args()

    today = dt.date.today()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    o.add_title_slide(prs, f"Prepared {today.isoformat()}",
                      "Eastern Pacific Storms · maracoos_007 · Typhoon Choi-wan")

    # ── Eastern Pacific ────────────────────────────────────────────────
    section(prs, "Eastern Pacific / Mexico",
            "Hurricane Polo cold wake · Hurricane Rachel · 90% disturbance  ·  Temperature 0 m & OHC")

    print("\nPolo")
    section(prs, "Hurricane Polo", "Cold wake as Polo tracked toward Baja  ·  Sep 21–27")
    for vd in ("temperature_0m", "ocean_heat_content"):
        map_slides(prs, EP, vd, POLO_DAYS, args.hour, "Polo cold wake")

    print("\nRachel")
    section(prs, "Hurricane Rachel",
            "Followed Polo's path, then crossed warm water along western Mexico  ·  formed Sep 25")
    nhc_slides(prs, "Rachel", ["cone", "key_messages"], today.isoformat())
    for vd in ("temperature_0m", "ocean_heat_content"):
        map_slides(prs, EP, vd, RACHEL_DAYS + [today], args.hour, "Rachel")

    print("\nNew disturbance")
    section(prs, "Next System — 90% Formation Chance",
            "NHC 7-day outlook: low south of Tehuantepec, TD expected mid–late week")
    buf = o.fetch_image("https://www.nhc.noaa.gov/xgtwo/two_pac_7d1.png")
    if buf:
        o.add_image_slide(prs, "NHC Outlook", "Eastern Pacific 7-Day Tropical Weather Outlook",
                          [(buf, "NHC 7-Day Outlook")], today.isoformat(), "")
    for vd in ("temperature_0m", "ocean_heat_content"):
        map_slides(prs, EP, vd, [today], args.hour, "Latest")

    # ── maracoos_007 ───────────────────────────────────────────────────
    print("\nmaracoos_007")
    section(prs, "maracoos_007",
            f"Mid Atlantic Bight glider  ·  {GLIDER_START:%b %d} – {today:%b %d}  ·  includes Sep 25–27 nor'easter")
    glider_slides(prs, today.isoformat(), today)
    days = [GLIDER_START + dt.timedelta(days=i) for i in range((today - GLIDER_START).days + 1)]
    map_slides(prs, MAB, "temperature_0m", days, args.hour, "maracoos_007 region")

    # ── Choi-wan ───────────────────────────────────────────────────────
    print("\nChoi-wan")
    section(prs, "Typhoon Choi-wan",
            "Ocean Heat Content  ·  formed Sep 30 near 15.6°N 151.9°E → 27.3°N 145.9°E today (135 kt)")
    map_slides(prs, GUAM, "ocean_heat_content", CHOIWAN_DAYS, args.hour, "Choi-wan")

    prs.save(args.output)
    print(f"\nSaved {len(prs.slides)} slides -> {args.output}")


if __name__ == "__main__":
    main()
